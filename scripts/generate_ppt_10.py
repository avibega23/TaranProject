#!/usr/bin/env python3
"""Generate a 10-slide FindIt presentation (condensed from the full deck)."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Pt

OUT = (
    Path(__file__).resolve().parents[1]
    / "deliverables"
    / "FindIt_Minor_Project_Presentation_10_Slides.pptx"
)

W, H = Emu(12192000), Emu(6858000)
FONT = "Calibri"

PAPER = RGBColor(0xF7, 0xF6, 0xF2)
INK = RGBColor(0x17, 0x20, 0x33)
MUTED = RGBColor(0x59, 0x65, 0x79)
TEAL = RGBColor(0x1F, 0x6B, 0x5A)
PALE = RGBColor(0xE6, 0xF2, 0xEE)
CORAL = RGBColor(0xCF, 0x67, 0x4F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x17, 0x1B, 0x2C)
SOFT = RGBColor(0xA9, 0xAE, 0xC0)

MENTOR = "Dr. Harpreet Kaur"


def rgb_fill(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def no_line(shape) -> None:
    shape.line.fill.background()


def add_rect(slide, x, y, w, h, fill: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    rgb_fill(shape, fill)
    no_line(shape)
    return shape


def set_run(paragraph, text, size, color, bold=False):
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return run


def add_text(slide, text, x, y, w, h, size=22, color=INK, bold=False, align="left", anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    try:
        tf._txBody.bodyPr.set("anchor", "t" if anchor == MSO_ANCHOR.TOP else "ctr")
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
    p.clear()
    set_run(p, text, size, color, bold=bold)
    return box


def add_multiline(slide, lines, x, y, w, h, size=16, color=MUTED, bold=False, align="left", space=6):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
        p.space_after = Pt(space)
        set_run(p, line, size, color, bold=bold)
    return box


def foot(slide, n, dark=False):
    col = SOFT if dark else MUTED
    add_text(slide, "FindIt", Emu(685800), Emu(6426200), Emu(4572000), Emu(228600), 12, col, True)
    add_text(slide, str(n), Emu(11125200), Emu(6426200), Emu(457200), Emu(228600), 12, col, True, "right")


def blank(prs, fill: RGBColor):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, W, H, fill)
    return slide


def card(slide, x, y, w, h, fill=PALE):
    return add_rect(slide, x, y, w, h, fill)


def build() -> None:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # ----- 1 Cover -----
    s = blank(prs, DARK)
    add_rect(s, 0, 0, Emu(152400), H, TEAL)
    add_text(s, "PUNJABI UNIVERSITY, PATIALA", Emu(685800), Emu(457200), Emu(10000000), Emu(304800), 14, SOFT, True)
    add_text(
        s,
        "Department of Computer Science & Engineering",
        Emu(685800),
        Emu(762000),
        Emu(10000000),
        Emu(304800),
        16,
        SOFT,
    )
    add_text(s, "FindIt", Emu(685800), Emu(1676400), Emu(9000000), Emu(914400), 60, WHITE, True)
    add_text(
        s,
        "Smart campus lost & found — report, match, and reunite belongings",
        Emu(685800),
        Emu(2590800),
        Emu(9500000),
        Emu(457200),
        20,
        RGBColor(0xD2, 0xD5, 0xE0),
    )
    add_text(
        s,
        "MINOR PROJECT  |  WEB DEV / FULL STACK",
        Emu(685800),
        Emu(3352800),
        Emu(9000000),
        Emu(304800),
        14,
        RGBColor(0x9F, 0xD0, 0xC0),
        True,
    )
    presenters = "Presented by Taran Singh (12401067)  ·  Prince (12401062)"
    add_text(s, presenters, Emu(685800), Emu(4191000), Emu(9500000), Emu(304800), 18, WHITE)
    add_text(s, "B.Tech CSE", Emu(685800), Emu(4572000), Emu(5000000), Emu(304800), 16, SOFT)
    add_text(s, f"Under the guidance of {MENTOR}", Emu(685800), Emu(5143500), Emu(9000000), Emu(304800), 16, RGBColor(0x9F, 0xD0, 0xC0))
    foot(s, 1, dark=True)

    # ----- 2 Problem Statement -----
    s = blank(prs, PAPER)
    add_text(s, "Problem Statement", Emu(685800), Emu(457200), Emu(9000000), Emu(457200), 36, INK, True)
    card(s, Emu(685800), Emu(1371600), Emu(4953000), Emu(4200000), RGBColor(0xF7, 0xE2, 0xDB))
    card(s, Emu(6096000), Emu(1371600), Emu(4953000), Emu(4200000), PALE)
    add_text(s, "1", Emu(838200), Emu(1524000), Emu(600000), Emu(400000), 28, CORAL, True)
    add_text(s, "The Real-World Challenge", Emu(838200), Emu(1981200), Emu(4500000), Emu(400000), 20, INK, True)
    add_multiline(
        s,
        [
            "Lost items are announced in scattered chats and notice boards.",
            "Posts are hard to search and easy to miss.",
            "Private identifying details often end up in public threads.",
        ],
        Emu(838200),
        Emu(2590800),
        Emu(4500000),
        Emu(2500000),
        16,
        MUTED,
        space=10,
    )
    add_text(s, "2", Emu(6248400), Emu(1524000), Emu(600000), Emu(400000), 28, TEAL, True)
    add_text(s, "The Need for a Simple Tool", Emu(6248400), Emu(1981200), Emu(4500000), Emu(400000), 20, INK, True)
    add_multiline(
        s,
        [
            "One searchable place for lost and found reports.",
            "Explainable match suggestions between opposite-type reports.",
            "Claims and moderation that keep ownership proof private.",
        ],
        Emu(6248400),
        Emu(2590800),
        Emu(4500000),
        Emu(2500000),
        16,
        MUTED,
        space=10,
    )
    foot(s, 2)

    # ----- 3 Objectives + Approach -----
    s = blank(prs, PAPER)
    add_text(s, "Objectives & Approach", Emu(685800), Emu(400000), Emu(9000000), Emu(400000), 34, INK, True)
    objectives = [
        ("Prepare Trustworthy Records", "Moderate lost/found reports for campus search and recovery."),
        ("Learn Useful Similarities", "Score opposite-type reports with category, text, location, and date — and explain why."),
        ("Make It Accessible", "Ship a Next.js app to browse, report, claim, and administer."),
    ]
    xs = [Emu(685800), Emu(4191000), Emu(7696200)]
    for i, (title, body) in enumerate(objectives):
        card(s, xs[i], Emu(1000000), Emu(3200400), Emu(2800000), PALE)
        add_text(s, "✓", xs[i] + Emu(152400), Emu(1120000), Emu(600000), Emu(350000), 24, TEAL, True)
        add_text(s, title, xs[i] + Emu(152400), Emu(1550000), Emu(2800000), Emu(700000), 16, INK, True)
        add_text(s, body, xs[i] + Emu(152400), Emu(2300000), Emu(2800000), Emu(1200000), 13, MUTED)
    card(s, Emu(685800), Emu(4200000), Emu(10668000), Emu(1700000), TEAL)
    add_text(s, "FindIt Path", Emu(900000), Emu(4400000), Emu(4000000), Emu(350000), 18, WHITE, True)
    add_text(
        s,
        "Moderated listings + explainable matches + private claims — instead of plain public CRUD or moderation-only boards.",
        Emu(900000),
        Emu(4900000),
        Emu(10000000),
        Emu(700000),
        15,
        RGBColor(0xD2, 0xD5, 0xE0),
    )
    foot(s, 3)

    # ----- 4 Data Model & Match Signals -----
    s = blank(prs, PAPER)
    add_text(s, "Data Model & Match Signals", Emu(685800), Emu(350000), Emu(10000000), Emu(400000), 32, INK, True)
    entities = [
        ("User / Session", "Accounts & signed-in cookies"),
        ("Item", "Lost or found report"),
        ("Claim", "Ownership request on a found item"),
        ("ActivityLog", "Moderation & claim history"),
    ]
    for i, (t, b) in enumerate(entities):
        x = Emu(685800) + Emu(i * 2700000)
        card(s, x, Emu(900000), Emu(2500000), Emu(1400000), PALE)
        add_text(s, t, x + Emu(120000), Emu(1000000), Emu(2200000), Emu(450000), 15, TEAL, True)
        add_text(s, b, x + Emu(120000), Emu(1500000), Emu(2200000), Emu(600000), 13, MUTED)
    signals = [
        ("Category ≈ 35%", "Same class pairs are likelier"),
        ("Text ≈ 35%", "Shared keywords in title/description"),
        ("Location ≈ 15%", "Campus geography hint"),
        ("Date ≈ 15%", "Decays outside 14 days"),
    ]
    add_text(s, "Default match weights (min score 0.35)", Emu(685800), Emu(2500000), Emu(9000000), Emu(350000), 16, MUTED)
    for i, (t, b) in enumerate(signals):
        x = Emu(685800) + Emu(i * 2700000)
        card(s, x, Emu(3000000), Emu(2500000), Emu(2700000), PALE if i % 2 == 0 else RGBColor(0xDC, 0xEF, 0xE8))
        add_text(s, t, x + Emu(120000), Emu(3200000), Emu(2200000), Emu(500000), 16, TEAL, True)
        add_text(s, b, x + Emu(120000), Emu(3800000), Emu(2200000), Emu(1400000), 13, MUTED)
    add_text(
        s,
        "≈70% of the default weight sits on category + text overlap. Matching is a heuristic estimate, not a probability.",
        Emu(685800),
        Emu(5900000),
        Emu(10668000),
        Emu(400000),
        14,
        MUTED,
    )
    foot(s, 4)

    # ----- 5 Tech Stack & Architecture -----
    s = blank(prs, PAPER)
    add_text(s, "Technology & Architecture", Emu(685800), Emu(350000), Emu(10000000), Emu(400000), 32, INK, True)
    stack = [
        ("Next.js / React", "App Router, TypeScript, Tailwind"),
        ("Express", "REST API, controllers, services"),
        ("MongoDB / Mongoose", "Documents and indexes"),
        ("Zod / RHF", "Request and form validation"),
        ("Vitest / Supertest", "In-memory API tests"),
        ("HTTP-only cookie", "Session auth (no localStorage tokens)"),
    ]
    for i, (title, body) in enumerate(stack):
        col, row = i % 3, i // 3
        x = Emu(685800) + Emu(col * 3500000)
        y = Emu(900000) + Emu(row * 1600000)
        card(s, x, y, Emu(3300000), Emu(1400000), PALE)
        add_text(s, title, x + Emu(140000), y + Emu(200000), Emu(3000000), Emu(400000), 15, TEAL, True)
        add_text(s, body, x + Emu(140000), y + Emu(650000), Emu(3000000), Emu(550000), 13, MUTED)
    layers = "Presentation (Next.js)  →  API Layer (Express rules)  →  Data Layer (MongoDB)"
    card(s, Emu(685800), Emu(4400000), Emu(10668000), Emu(1500000), DARK)
    add_text(s, "Three separated layers", Emu(900000), Emu(4600000), Emu(9000000), Emu(350000), 14, RGBColor(0x9F, 0xD0, 0xC0), True)
    add_text(s, layers, Emu(900000), Emu(5100000), Emu(10000000), Emu(500000), 18, WHITE)
    foot(s, 5)

    # ----- 6 Live Demo -----
    s = blank(prs, PAPER)
    add_text(s, "Live App Demo", Emu(685800), Emu(350000), Emu(10000000), Emu(400000), 32, INK, True)
    card(s, Emu(685800), Emu(950000), Emu(5200000), Emu(4800000), PALE)
    add_text(s, "Browse & Report", Emu(900000), Emu(1150000), Emu(4600000), Emu(400000), 20, TEAL, True)
    add_multiline(
        s,
        [
            "Home shows approved lost / found / resolved counts.",
            "Browse filters by type, category, location, and date.",
            "Report form creates a pending item (optional image).",
            "Pending reports stay off public browse until approved.",
        ],
        Emu(900000),
        Emu(1700000),
        Emu(4600000),
        Emu(3500000),
        15,
        INK,
        space=10,
    )
    card(s, Emu(6400000), Emu(950000), Emu(5000000), Emu(4800000), DARK)
    add_text(s, "Match Explanation", Emu(6600000), Emu(1150000), Emu(4600000), Emu(400000), 20, RGBColor(0x9F, 0xD0, 0xC0), True)
    add_multiline(
        s,
        [
            "Open a lost earbuds report from seed data.",
            "Ranked found candidates with scores and reasons:",
            "  • Same category  • Shared words",
            "  • Similar location  • Days apart",
            "Reviewers see why reports were linked before approving claims.",
            "Disclaimer: heuristic estimate, not a probability.",
        ],
        Emu(6600000),
        Emu(1700000),
        Emu(4600000),
        Emu(3500000),
        14,
        RGBColor(0xD2, 0xD5, 0xE0),
        space=8,
    )
    foot(s, 6)

    # ----- 7 Sample Walkthrough -----
    s = blank(prs, PAPER)
    add_text(s, "Sample Walkthrough", Emu(685800), Emu(400000), Emu(9000000), Emu(400000), 32, INK, True)
    add_text(
        s,
        "The exact seeded scenario used in the demo, and what the system does",
        Emu(685800),
        Emu(900000),
        Emu(10000000),
        Emu(300000),
        16,
        MUTED,
    )
    steps = [
        ("Seed", "Demo students + earbuds reports"),
        ("Match", "Lost ↔ found with reasons"),
        ("Claim", "Proof visible to reviewers only"),
        ("Resolve", "Item marked resolved"),
    ]
    xs = [Emu(685800), Emu(3429000), Emu(6172200), Emu(8915400)]
    for i, (t, b) in enumerate(steps):
        card(s, xs[i], Emu(1500000), Emu(2500000), Emu(2200000), PALE if i % 2 == 0 else RGBColor(0xDC, 0xEF, 0xE8))
        add_text(s, t, xs[i] + Emu(120000), Emu(1700000), Emu(2200000), Emu(500000), 20, TEAL, True)
        add_text(s, b, xs[i] + Emu(120000), Emu(2400000), Emu(2200000), Emu(1000000), 14, MUTED)
    card(s, Emu(685800), Emu(4200000), Emu(10668000), Emu(1500000), DARK)
    add_text(s, "Predicted outcome", Emu(900000), Emu(4400000), Emu(4000000), Emu(350000), 14, RGBColor(0x9F, 0xD0, 0xC0), True)
    add_text(
        s,
        "A recoverable campus loop: searchable listings, explainable matches, moderated claims — not a chat dump.",
        Emu(900000),
        Emu(4850000),
        Emu(10000000),
        Emu(700000),
        16,
        WHITE,
    )
    foot(s, 7)

    # ----- 8 Validation -----
    s = blank(prs, PAPER)
    add_text(s, "Validation & Demo Success", Emu(685800), Emu(400000), Emu(10000000), Emu(400000), 32, INK, True)
    add_text(
        s,
        "Checked with automated API tests and a seeded walkthrough — not ML hold-out metrics",
        Emu(685800),
        Emu(900000),
        Emu(10000000),
        Emu(350000),
        15,
        MUTED,
    )
    metrics = [
        ("Vitest", "API behaviours covered with Supertest + in-memory MongoDB"),
        ("Privacy", "Public JSON omits identifying details and claim proof"),
        ("Demo path", "Seed → match → claim → resolve works end-to-end"),
    ]
    for i, (k, v) in enumerate(metrics):
        y = Emu(1500000) + Emu(i * 1300000)
        card(s, Emu(685800), y, Emu(10668000), Emu(1100000), PALE)
        add_text(s, k, Emu(900000), y + Emu(250000), Emu(2500000), Emu(400000), 22, TEAL, True)
        add_text(s, v, Emu(3600000), y + Emu(300000), Emu(7000000), Emu(500000), 18, INK)
    foot(s, 8)

    # ----- 9 Conclusion -----
    s = blank(prs, PAPER)
    add_text(s, "Conclusion & Future Work", Emu(685800), Emu(400000), Emu(10000000), Emu(400000), 32, INK, True)
    card(s, Emu(685800), Emu(1100000), Emu(5200000), Emu(4500000), PALE)
    card(s, Emu(6200000), Emu(1100000), Emu(5200000), Emu(4500000), RGBColor(0xF7, 0xE2, 0xDB))
    add_text(s, "What We Achieved", Emu(900000), Emu(1300000), Emu(4600000), Emu(400000), 20, TEAL, True)
    add_multiline(
        s,
        [
            "• Full-stack FindIt with Next.js + Express + MongoDB",
            "• Explainable match weights (35/35/15/15)",
            "• Private claim proof + admin moderation",
            "• Seeded demo path and automated tests",
        ],
        Emu(900000),
        Emu(1900000),
        Emu(4600000),
        Emu(3200000),
        15,
        INK,
        space=10,
    )
    add_text(s, "Future Work", Emu(6400000), Emu(1300000), Emu(4600000), Emu(400000), 20, CORAL, True)
    add_multiline(
        s,
        [
            "• Email / in-app notifications on claim review",
            "• Map or building list for locations",
            "• Object storage for images",
            "• Richer administrator audit views",
        ],
        Emu(6400000),
        Emu(1900000),
        Emu(4600000),
        Emu(3200000),
        15,
        INK,
        space=10,
    )
    foot(s, 9)

    # ----- 10 Thank you -----
    s = blank(prs, DARK)
    add_rect(s, 0, 0, Emu(152400), H, TEAL)
    add_text(s, "Thank You", Emu(685800), Emu(2000000), Emu(10000000), Emu(800000), 54, WHITE, True)
    add_text(s, "Questions are welcome", Emu(685800), Emu(2900000), Emu(9000000), Emu(400000), 22, SOFT)
    add_text(
        s,
        "Taran Singh | 12401067    ·    Prince | 12401062",
        Emu(685800),
        Emu(4000000),
        Emu(10000000),
        Emu(400000),
        18,
        WHITE,
    )
    add_text(s, "B.Tech CSE  |  Punjabi University, Patiala", Emu(685800), Emu(4500000), Emu(9000000), Emu(350000), 16, SOFT)
    add_text(s, f"Under the guidance of {MENTOR}", Emu(685800), Emu(5100000), Emu(9000000), Emu(350000), 16, RGBColor(0x9F, 0xD0, 0xC0))
    foot(s, 10, dark=True)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"Wrote {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
